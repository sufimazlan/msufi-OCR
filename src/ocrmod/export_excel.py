"""Step 8: Excel export with colour flags. One row per document, one per line item, and
a Flags sheet listing every yellow/red field with its reason and crop."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .schemas import DocumentResult

FILL = {
    "red": PatternFill("solid", fgColor="F8CBAD"),
    "yellow": PatternFill("solid", fgColor="FFE699"),
}
HEADER_FONT = Font(bold=True)

DOC_COLUMNS = [
    ("source_name", "File"), ("company", "Company"), ("doc_type", "Doc type"),
    ("status", "Status"), ("needs_review", "Needs review"), ("vendor_code", "Vendor code"),
    ("vendor_name", "Vendor"), ("vendor_tax_id", "Vendor tax ID"),
    ("invoice_number", "Invoice no."), ("invoice_date", "Invoice date"), ("due_date", "Due date"),
    ("po_number", "PO no."), ("currency", "Currency"), ("subtotal", "Subtotal"),
    ("tax_rate", "Tax %"), ("tax_amount", "Tax"), ("total", "Total"),
]
DOC_LEVEL = {"source_name", "company", "doc_type", "status", "needs_review", "vendor_code"}
LINE_COLUMNS = ["description", "quantity", "unit_price", "amount"]


def _header(ws, names):
    ws.append(names)
    for cell in ws[1]:
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"


def _flag(cell, fr):
    if fr is not None and fr.level in FILL:
        cell.fill = FILL[fr.level]
        if fr.reasons:
            cell.comment = Comment("; ".join(fr.reasons), "ocrmod")


def _autosize(ws):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 8), 50)


def export(results: list[DocumentResult], path: Path) -> Path:
    wb = Workbook()
    docs = wb.active
    docs.title = "Documents"
    _header(docs, [label for _, label in DOC_COLUMNS])
    lines = wb.create_sheet("Line items")
    _header(lines, ["File", "Line", *[c.replace("_", " ").title() for c in LINE_COLUMNS]])
    flags = wb.create_sheet("Flags")
    _header(flags, ["File", "Field", "Value", "Flag", "Reasons", "Score", "Page", "Crop"])

    for res in results:
        row = []
        for key, _ in DOC_COLUMNS:
            if key in DOC_LEVEL:
                v = getattr(res, key)
                row.append("yes" if v is True else "no" if v is False else v)
            else:
                row.append(res.extraction.get(key))
        docs.append(row)
        r = docs.max_row
        for col, (key, _) in enumerate(DOC_COLUMNS, start=1):
            if key not in DOC_LEVEL:
                _flag(docs.cell(r, col), res.field(key))
        if res.status == "failed":
            docs.cell(r, 4).fill = FILL["red"]
            docs.cell(r, 4).comment = Comment(res.error or "failed", "ocrmod")

        for n, item in enumerate(res.extraction.get("line_items") or []):
            lines.append([res.source_name, n + 1, *[item.get(c) for c in LINE_COLUMNS]])
            lr = lines.max_row
            for col, c in enumerate(LINE_COLUMNS, start=3):
                _flag(lines.cell(lr, col), res.field(f"line_items.{n}.{c}"))

        for f in res.fields:
            if f.level == "green":
                continue
            crop = f"{res.folder}/{f.crop}" if f.crop else None
            flags.append([res.source_name, f.path, f.value, f.level.upper(), "; ".join(f.reasons),
                          f.score, None if f.page is None else f.page + 1, crop])
            flags.cell(flags.max_row, 4).fill = FILL[f.level]

    about = wb.create_sheet("About")
    calibrated = all(r.thresholds_calibrated for r in results) if results else False
    about.append(["Thresholds calibrated", "yes" if calibrated else
                  "NO - placeholder thresholds, flags are indicative only"])
    about.append(["Red", "Validation failed or very low confidence: must check"])
    about.append(["Yellow", "Softer check failed or lower confidence: please look"])
    about.append(["Comments", "Hover a coloured cell to see the reason"])

    for ws in (docs, lines, flags, about):
        _autosize(ws)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
