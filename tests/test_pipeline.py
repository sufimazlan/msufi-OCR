from openpyxl import load_workbook

from conftest import SAMPLES, FakeExtractor
from ocrmod.config import Settings
from ocrmod.export_excel import export
from ocrmod.pipeline import load_results, run_batch


def test_pipeline_end_to_end_on_samples(tmp_path):
    settings = Settings.load()
    settings.data["ingest"]["use_paddle_orientation"] = False
    pdfs = sorted(SAMPLES.glob("*.pdf"))
    results = run_batch(pdfs, tmp_path / "run", settings, company="SUB-A", layout_name="pdftext",
                        lines_name="pdftext", extractor=FakeExtractor(),
                        registry_path=tmp_path / "registry.sqlite")
    by = {r.source_name: r for r in results}

    clean = by["inv-01-layout-a.pdf"]
    assert clean.status == "done" and clean.vendor_code == "V0001"
    assert not clean.needs_review, [(f.path, f.reasons) for f in clean.fields if f.level != "green"]

    assert by["inv-02-layout-b.pdf"].vendor_code == "V0002"

    sub = by["inv-03-layout-c.pdf"].field("subtotal")
    assert sub.level == "red" and "sum mismatch" in sub.reasons

    dup = by["inv-04-layout-a.pdf"].field("invoice_number")
    assert dup.level == "red" and "duplicate invoice" in dup.reasons

    vendor = by["inv-05-layout-b.pdf"].field("vendor_name")
    assert vendor.level == "yellow" and "vendor not in master" in vendor.reasons

    # Image-only scan has no text layer, so the dev engine finds nothing: all required red.
    scan = by["scan-01-skewed.pdf"]
    assert scan.needs_review and scan.field("total").level == "red"

    # Outputs on disk, and a re-run of the same file is not reported as a duplicate.
    assert len(load_results(tmp_path / "run")) == len(pdfs)
    again = run_batch([SAMPLES / "inv-02-layout-b.pdf"], tmp_path / "run2", settings, layout_name="pdftext",
                      lines_name="pdftext", extractor=FakeExtractor(),
                      registry_path=tmp_path / "registry.sqlite")
    assert not again[0].needs_review

    xlsx = export(results, tmp_path / "out.xlsx")
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Documents", "Line items", "Flags", "About"]
    assert wb["Documents"].max_row == len(pdfs) + 1
    assert wb["Flags"].max_row > 1


def test_low_token_probability_flags_field(tmp_path):
    settings = Settings.load()
    settings.data["ingest"]["use_paddle_orientation"] = False
    pdf = SAMPLES / "inv-01-layout-a.pdf"
    res = run_batch([pdf], tmp_path / "run", settings, layout_name="pdftext", lines_name="pdftext",
                    extractor=FakeExtractor({"total": 0.5}),
                    registry_path=tmp_path / "r.sqlite")[0]
    total = res.field("total")
    assert total.level == "red" and total.crop and (tmp_path / "run" / res.folder / total.crop).exists()
