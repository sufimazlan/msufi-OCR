from fastapi.testclient import TestClient
from openpyxl import load_workbook

from conftest import SAMPLES, FakeExtractor
from ocrmod.config import Settings
from ocrmod.web.app import create_app
from ocrmod.web.db import Batch, Document


def make_client(tmp_path):
    settings = Settings.load()
    settings.data["paths"]["data_dir"] = str(tmp_path / "data")
    settings.data["paths"]["registry"] = str(tmp_path / "data" / "registry.sqlite")
    settings.data["ingest"]["use_paddle_orientation"] = False
    app = create_app(settings, start_worker=False, extractor=FakeExtractor(),
                     layout_name="pdftext", lines_name="pdftext")
    return TestClient(app), app


def upload(client, *names, company="SUB-A"):
    files = [("files", (n, (SAMPLES / n).read_bytes(), "application/pdf")) for n in names]
    return client.post("/batches", files=files, data={"company": company, "doc_type": "invoice"},
                       follow_redirects=False)


def test_upload_progress_and_excel_download(tmp_path):
    client, app = make_client(tmp_path)
    with client:
        r = upload(client, "inv-01-layout-a.pdf", "inv-03-layout-c.pdf", "inv-05-layout-b.pdf")
        assert r.status_code == 303 and r.headers["location"] == "/batches/1"

        # Queued: page polls, no download yet; files saved in the data folder, not the repo.
        status = client.get("/batches/1/status").text
        assert "Waiting in the queue" in status and 'hx-trigger="every' in status
        assert client.get("/batches/1/export.xlsx").status_code == 404
        assert len(list((tmp_path / "data" / "batches" / "00001" / "input").rglob("*.pdf"))) == 3

        worker = app.state.worker
        assert worker.process_next() is True
        assert worker.process_next() is False  # queue empty
        with app.state.sessions() as s:
            batch = s.get(Batch, 1)
            assert batch.status == "done"
            docs = {d.filename: d for d in batch.documents}
        assert docs["inv-01-layout-a.pdf"].flags == 0
        assert docs["inv-03-layout-c.pdf"].red_flags >= 1      # sum mismatch
        assert docs["inv-05-layout-b.pdf"].flags >= 1          # vendor not in master

        status = client.get("/batches/1/status").text
        assert "Download Excel" in status and "hx-trigger" not in status  # polling stops
        assert "to check" in status and "clean" in status

        r = client.get("/batches/1/export.xlsx")
        assert r.status_code == 200 and "ocr-SUB-A-1-" in r.headers["content-disposition"]
        xlsx = tmp_path / "dl.xlsx"
        xlsx.write_bytes(r.content)
        sheet = load_workbook(xlsx)["Documents"]
        assert sheet.max_row == 4
        assert [sheet.cell(r, 1).value for r in (2, 3, 4)] == [
            "inv-01-layout-a.pdf", "inv-03-layout-c.pdf", "inv-05-layout-b.pdf"]

        home = client.get("/").text
        assert "Recent batches" in home and "/batches/1/export.xlsx" in home


def test_progress_stages_reported_in_order(tmp_path):
    from ocrmod.pipeline import run_batch

    settings = Settings.load()
    settings.data["ingest"]["use_paddle_orientation"] = False
    events = []
    run_batch([SAMPLES / "inv-01-layout-a.pdf", SAMPLES / "inv-02-layout-b.pdf"], tmp_path / "out",
              settings, layout_name="pdftext", lines_name="pdftext", extractor=FakeExtractor(),
              registry_path=tmp_path / "r.sqlite", progress=lambda i, st: events.append((i, st)))
    assert events == [(0, "ingest"), (1, "ingest"), (0, "layout"), (1, "layout"),
                      (0, "lines"), (1, "lines"), (0, "extract"), (0, "done"),
                      (1, "extract"), (1, "done")]


def test_rejects_non_pdf_and_leaves_nothing_behind(tmp_path):
    client, app = make_client(tmp_path)
    with client:
        r = client.post("/batches", files=[("files", ("notes.pdf", b"hello", "application/pdf"))],
                        follow_redirects=False)
        assert r.status_code == 400 and "not a PDF" in r.text
        with app.state.sessions() as s:
            assert s.query(Batch).count() == 0


def test_restart_requeues_interrupted_batch(tmp_path):
    client, app = make_client(tmp_path)
    with client:
        upload(client, "inv-02-layout-b.pdf")
        with app.state.sessions.begin() as s:
            b = s.get(Batch, 1)
            b.status = "processing"
            b.documents[0].status, b.documents[0].stage = "processing", "layout"
        app.state.worker.recover()
        with app.state.sessions() as s:
            b = s.get(Batch, 1)
            assert b.status == "queued" and b.documents[0].stage is None
        assert app.state.worker.process_next() is True
        with app.state.sessions() as s:
            assert s.get(Document, 1).status == "done"


def test_batch_failure_is_shown(tmp_path):
    client, app = make_client(tmp_path)

    class Broken(FakeExtractor):
        def check(self):
            raise RuntimeError("Cannot reach Ollama at http://localhost:11434")

    app.state.worker.extractor = Broken()
    with client:
        upload(client, "inv-01-layout-a.pdf")
        app.state.worker.process_next()
        status = client.get("/batches/1/status").text
        assert "Batch failed" in status and "Cannot reach Ollama" in status
        assert "hx-trigger" not in status


def test_same_file_name_twice_in_one_batch(tmp_path):
    client, app = make_client(tmp_path)
    with client:
        pdf = (SAMPLES / "inv-02-layout-b.pdf").read_bytes()
        files = [("files", ("scan.pdf", pdf, "application/pdf")),
                 ("files", ("scan.pdf", pdf, "application/pdf"))]
        assert client.post("/batches", files=files, follow_redirects=False).status_code == 303
        app.state.worker.process_next()
        with app.state.sessions() as s:
            assert [d.status for d in s.get(Batch, 1).documents] == ["done", "done"]
